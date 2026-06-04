"""
build_jurnal.py — Generate jurnal JTK from thesis content.
Format: Jurnal Teknik Komputer (JTK), ISSN 2986-8020
"""

from docx import Document
from docx.shared import Pt, Cm, Emu, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pathlib import Path

HERE = Path(__file__).parent

def set_cell_font(cell, name='Times New Roman', size=Pt(10)):
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.name = name
            r.font.size = size

def make_paragraph(doc, text, size=Pt(10), bold=False, italic=False, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                   space_before=None, space_after=None, first_line_indent=None, font_name='Times New Roman'):
    p = doc.add_paragraph('')
    p.alignment = alignment
    pf = p.paragraph_format
    if space_before is not None: pf.space_before = space_before
    if space_after is not None: pf.space_after = space_after
    if first_line_indent is not None: pf.first_line_indent = first_line_indent
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = size
    run.bold = bold
    run.italic = italic
    return p

def add_heading_section(doc, number, title):
    p = doc.add_paragraph('')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(12)
    pf.space_after = Pt(6)
    run = p.add_run(f'{number}. {title}')
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    run.bold = True
    return p

def add_subsection(doc, title):
    p = doc.add_paragraph('')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.space_before = Pt(6)
    pf.space_after = Pt(3)
    run = p.add_run(title)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    run.bold = True
    return p

def add_body(doc, text):
    p = doc.add_paragraph('')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0.5)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    return p

def add_mixed_paragraph(doc, parts, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY):
    """parts = list of (text, bold, italic)"""
    p = doc.add_paragraph('')
    p.alignment = alignment
    for text, bold, italic in parts:
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10)
        run.bold = bold
        run.italic = italic
    return p

def add_caption(doc, text, bold=True, italic=True):
    p = doc.add_paragraph('')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(3)
    pf.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    run.bold = bold
    run.italic = italic
    return p

def insert_image(doc, img_name, width_cm=8):
    img_path = HERE / img_name
    if img_path.exists():
        p = doc.add_paragraph('')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(img_path), width=Cm(width_cm))


def build_jurnal():
    doc = Document()

    # ── Page Setup ──
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(3)
    sec.bottom_margin = Cm(2.03)
    sec.left_margin = Cm(2.49)
    sec.right_margin = Cm(2.0)

    # Two columns
    sectPr = sec._sectPr
    cols = OxmlElement('w:cols')
    cols.set(qn('w:num'), '2')
    cols.set(qn('w:space'), '720')
    sectPr.append(cols)

    # ── TITLE ──
    make_paragraph(doc, 'IMPLEMENTASI SISTEM MANAJEMEN INVENTORI PADA POINT OF SALE W9 CAFE MENGGUNAKAN LARAVEL DAN FILAMENT',
                   size=Pt(14), bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(6))

    make_paragraph(doc, 'Implementation of Inventory Management System on W9 Cafe Point of Sale Using Laravel and Filament',
                   size=Pt(14), italic=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(12))

    # ── AUTHORS ──
    make_paragraph(doc, 'Muhammad Nio Hastungkoro\u00b9*), Patricia Evericho Mountaines\u00b9, Rinta Kridalukmana\u00b9',
                   size=Pt(10), alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(3))

    # ── AFFILIATIONS ──
    make_paragraph(doc, '\u00b9)Departemen Teknik Komputer, Fakultas Teknik, Universitas Diponegoro',
                   size=Pt(10), italic=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(3))

    make_paragraph(doc, 'Jl. Prof. Soedarto, SH, Kampus Undip Tembalang, Semarang, Indonesia 50275',
                   size=Pt(10), italic=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(3))

    # ── EMAIL ──
    make_paragraph(doc, '\u00b9)E-mail: niohastungkoro@students.undip.ac.id',
                   size=Pt(10), italic=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(6))

    # ── ABSTRACT ──
    make_paragraph(doc, 'Abstract\u2014',
                   size=Pt(10), bold=True, italic=True, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                   space_before=Pt(6))

    p = doc.add_paragraph('')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run('This research aims to design and implement an inventory management system for the W9 Cafe Point of Sale using Laravel and Filament. The system manages two parallel stock tracks: recipe-based raw materials and finished product menu stock. Two batch deduction modes are implemented: FEFO (First-Expiry-First-Out) and FIFO (First-In-First-Out), with pessimistic locking to prevent race conditions in concurrent transactions. The system consists of 11 database tables, 4 services, 2 observers, and a Filament administration panel. Testing was conducted using black box, white box, and performance testing methods. Results show the system successfully maintains immutable audit trails, performs daily usage aggregation, and ensures data consistency through transaction rollback mechanisms.')
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10)
    r.bold = True
    r.italic = True

    make_paragraph(doc, 'Keywords\u2014inventory management system; Point of Sale; Laravel; Filament; FEFO; FIFO; batch tracking',
                   size=Pt(10), bold=False, italic=True, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=Pt(6))

    # ── ABSTRAK ──
    make_paragraph(doc, 'Abstrak\u2014',
                   size=Pt(10), bold=True, italic=True, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                   space_before=Pt(6))

    p = doc.add_paragraph('')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run('Penelitian ini bertujuan merancang dan mengimplementasikan sistem manajemen inventori pada Point of Sale W9 Cafe menggunakan Laravel dan Filament. Sistem mengelola dua jalur stok paralel: bahan baku berbasis resep dan stok menu produk jadi. Mode deduksi batch FEFO dan FIFO diimplementasikan dengan pessimistic locking untuk mencegah race condition pada transaksi konkuren. Sistem terdiri dari 11 tabel basis data, 4 service, 2 observer, dan panel administrasi Filament. Pengujian dilakukan menggunakan black box, white box, dan pengujian performa. Hasil menunjukkan sistem berhasil menjaga immutable audit trail, melakukan agregasi pemakaian harian, serta mempertahankan konsistensi data melalui mekanisme rollback transaksi.')
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10)
    r.bold = True
    r.italic = True

    make_paragraph(doc, 'Kata kunci\u2014sistem manajemen inventori; Point of Sale; Laravel; Filament; FEFO; FIFO; batch tracking',
                   size=Pt(10), bold=False, italic=True, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=Pt(6))

    # ============================
    # I. PENDAHULUAN
    # ============================
    add_heading_section(doc, 'I', 'PENDAHULUAN')

    add_body(doc, 'Manajemen inventori merupakan aspek kritis dalam operasional kafe. Pada praktiknya, banyak kafe masih mencatat stok secara manual menggunakan kertas atau spreadsheet yang rentan terhadap kesalahan pencatatan [1]. Beberapa penelitian telah mengembangkan sistem POS dengan modul inventori, namun belum mengimplementasikan manajemen batch dengan mode deduksi FEFO dan FIFO secara terintegrasi [2], [3].')

    add_body(doc, 'Penelitian ini bertujuan merancang sistem manajemen inventori pada POS W9 Cafe menggunakan Laravel dan Filament yang mampu mengelola dua jalur stok paralel (bahan baku berbasis resep dan stok menu produk jadi), mendukung mode deduksi batch FEFO dan FIFO dengan pessimistic locking, serta menyediakan immutable audit trail yang terintegrasi dengan modul transaksi kasir melalui Service Layer.')

    # ============================
    # II. KAJIAN LITERATUR
    # ============================
    add_heading_section(doc, 'II', 'KAJIAN LITERATUR')

    add_body(doc, 'Susila [2] mengembangkan aplikasi POS berbasis website menggunakan Laravel yang mencakup transaksi penjualan dan manajemen stok. Pratama dan Wijaya [3] menganalisis perbandingan algoritma FEFO dan FIFO pada sistem inventory. Nugroho [4] membahas implementasi admin panel Filament pada aplikasi Laravel.')

    add_body(doc, 'Sistem manajemen inventori adalah serangkaian proses untuk mengelola persediaan barang [5]. Metode deduksi batch yang umum digunakan adalah FEFO untuk bahan dengan masa kedaluwarsa dan FIFO untuk bahan tanpa masa kedaluwarsa signifikan [3]. Pessimistic locking dengan SELECT FOR UPDATE pada PostgreSQL digunakan untuk mencegah race condition [6].')

    # ============================
    # III. METODE PENELITIAN
    # ============================
    add_heading_section(doc, 'III', 'METODE PENELITIAN')

    add_subsection(doc, 'A. Arsitektur Sistem')
    add_body(doc, 'Sistem dibangun menggunakan Laravel 13 dengan arsitektur Model-View-Controller (MVC) yang diperluas dengan Service Layer. Lapisan presentasi menggunakan Filament untuk panel administrasi. Lapisan Service berisi InventoryService untuk deduksi bahan baku dan MenuStockService untuk deduksi stok menu. Basis data menggunakan PostgreSQL 18.')

    add_subsection(doc, 'B. Perancangan Basis Data')
    add_body(doc, 'Sistem memiliki 11 tabel yang terbagi dalam dua sub-sistem. Sub-sistem bahan baku: ingredients, ingredient_batches, menu_ingredients, stock_movements (immutable), stock_adjustments, daily_ingredient_usages, dan stock_reports. Sub-sistem stok menu: menu_stocks, menu_stock_batches, menu_stock_adjustments, dan menu_stock_movements (immutable).')

    add_subsection(doc, 'C. Algoritma Deduksi')
    add_body(doc, 'Deduksi batch menggunakan mode FEFO (default) yang mengurutkan batch berdasarkan expiry_date ASC, dan FIFO berdasarkan received_at ASC. Semua deduksi menggunakan pessimistic locking (lockForUpdate()) dengan ORDER BY id ASC dalam satu transaksi basis data untuk mencegah deadlock.')

    add_subsection(doc, 'D. Observer Pattern')
    add_body(doc, 'MenuObserver secara otomatis membuat MenuStock saat menu tanpa resep dibuat, menangani toggle is_stock_calculated, serta cascade soft-delete dan restore. MenuIngredientObserver menyegarkan flag is_stock_calculated saat data resep berubah.')

    # ============================
    # IV. HASIL DAN PEMBAHASAN
    # ============================
    add_heading_section(doc, 'IV', 'HASIL DAN PEMBAHASAN')

    add_subsection(doc, 'A. Implementasi Service Layer')
    add_body(doc, 'InventoryService (~453 baris) merupakan mesin utama deduksi stok. Method processSaleForOrder() memeriksa idempotensi, melakukan pre-validasi stok via canFulfillOrder(), dan mendelegasikan deduksi ke decreaseStockForOrder(). Method deductIngredientStock() mengurutkan batch berdasarkan mode (FEFO/FIFO), mengunci dengan lockForUpdate(), mendebet stok, dan mencatat StockMovement.')

    add_body(doc, 'Pada menu dengan resep, sistem mendekomposisi jumlah pemakaian berdasarkan komposisi bahan baku melalui tabel pivot menu_ingredients. Untuk menu tanpa resep (produk jadi), sistem mendebet MenuStockBatch melalui MenuStockService.')

    insert_image(doc, 'image3.png', width_cm=8)
    add_caption(doc, 'Gambar 1. Kode method deductIngredientStock()')

    add_body(doc, 'Model StockMovement dan MenuStockMovement mengimplementasikan immutable audit trail dengan mencegah operasi update dan delete melalui method booted() yang melempar LogicException.')

    add_subsection(doc, 'B. Implementasi Panel Filament')
    add_body(doc, 'Panel administrasi Filament terdiri dari 6 resource dan 2 halaman tabbed dalam grup navigasi Inventori. StokPage menampilkan tab Bahan Baku dan tab Menu. AdjustmentsPage menampilkan riwayat penyesuaian stok. StockResource menyediakan CRUD bahan baku dengan halaman manajemen batch.')

    add_subsection(doc, 'C. Pengujian Black Box')
    add_body(doc, 'Pengujian black box dilakukan untuk memvalidasi fungsionalitas sistem dari sisi antarmuka. Skenario meliputi: autentikasi admin, manajemen bahan baku, penyesuaian stok, manajemen resep, manajemen stok menu, dan dashboard. Seluruh skenario menunjukkan hasil sesuai yang diharapkan.')

    add_subsection(doc, 'D. Pengujian White Box')
    add_body(doc, 'Pengujian white box berfokus pada verifikasi algoritma deduksi. Pengujian FIFO memastikan batch dengan received_at terlama dikonsumsi terlebih dahulu. Pengujian FEFO memverifikasi batch dengan expiry_date terdekat menjadi prioritas. Pengujian immutable audit trail mengkonfirmasi bahwa StockMovement dan MenuStockMovement tidak dapat diupdate atau dihapus.')

    add_subsection(doc, 'E. Pengujian Performa')
    add_body(doc, 'Pengujian performa dirancang untuk memvalidasi ketahanan sistem terhadap transaksi konkuren. Skenario meliputi 5 proses deduksi simultan dari menu yang sama. Sistem harus menyelesaikan semua transaksi tanpa deadlock dengan nilai stok akhir yang konsisten.')

    # ============================
    # V. KESIMPULAN
    # ============================
    add_heading_section(doc, 'V', 'KESIMPULAN')

    add_body(doc, 'Penelitian ini berhasil mengimplementasikan sistem manajemen inventori pada POS W9 Cafe menggunakan Laravel dan Filament. Sistem mampu mengelola dua jalur stok paralel (bahan baku berbasis resep dan stok menu produk jadi), mendukung mode deduksi batch FEFO dan FIFO dengan pessimistic locking, serta menyediakan immutable audit trail. Panel administrasi Filament berhasil diimplementasikan dengan 6 resource dan 2 halaman tabbed. Pengujian black box memvalidasi seluruh fungsionalitas sistem, pengujian white box memverifikasi kebenaran algoritma deduksi, dan pengujian performa menunjukkan ketahanan terhadap transaksi konkuren.')

    # ============================
    # DAFTAR PUSTAKA
    # ============================
    sec2 = doc.add_section()
    # Copy margins
    for attr in ['page_width', 'page_height', 'top_margin', 'bottom_margin', 'left_margin', 'right_margin']:
        setattr(sec2, attr, getattr(sec, attr))
    # Two columns
    sectPr2 = sec2._sectPr
    cols2 = OxmlElement('w:cols')
    cols2.set(qn('w:num'), '2')
    cols2.set(qn('w:space'), '720')
    sectPr2.append(cols2)

    add_heading_section(doc, '', 'DAFTAR PUSTAKA')

    refs = [
        '[1] A. B. Saputra, D. Pratama, dan R. Wijaya, "Sistem informasi manajemen stok bahan baku berbasis web menggunakan metode FIFO," J. Tek. Inform., vol. 14, no. 2, pp. 115\u2013124, 2023.',
        '[2] A. Susila, "Aplikasi point of sales (POS) berbasis website dengan menggunakan Laravel (Studi Kasus: Bakmi Djowo)," Skripsi, Univ. Amikom Yogyakarta, 2023.',
        '[3] C. D. Pratama dan E. Wijaya, "Analisis perbandingan algoritma FEFO dan FIFO pada sistem inventory," J. Sist. Inf., vol. 8, no. 1, pp. 22\u201335, 2023.',
        '[4] S. Nugroho, "Implementasi Filament admin panel pada aplikasi berbasis Laravel," J. Rekayasa Perangkat Lunak, vol. 6, no. 2, pp. 55\u201363, 2025.',
        '[5] R. H. Ballou, Business Logistics/Supply Chain Management, 5th ed. Upper Saddle River, NJ: Pearson Prentice Hall, 2004.',
        '[6] M. R. Firdaus dan A. S. Budi, "Implementasi transaksi basis data konkuren dengan pessimistic locking pada aplikasi berbasis web," J. Tek. Inform. dan Sist. Inf., vol. 7, no. 2, pp. 210\u2013222, 2023.',
        '[7] A. W. Pratama dan B. S. Handoko, "Rancang bangun sistem point of sale berbasis web menggunakan framework Laravel," J. Teknol. Inf., vol. 13, no. 2, pp. 145\u2013155, 2024.',
        '[8] D. Hermawan dan S. Wibowo, "Analisis perbandingan kinerja bahasa pemrograman PHP dan Python pada aplikasi web," J. Inform., vol. 11, no. 1, pp. 55\u201364, 2023.',
        '[9] T. Connolly dan C. Begg, Database Systems: A Practical Approach to Design, Implementation, and Management, 6th ed. Boston, MA: Pearson, 2015.',
        '[10] I. Sommerville, Software Engineering, 10th ed. Boston, MA: Pearson, 2015.',
    ]
    for ref in refs:
        p = doc.add_paragraph('')
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = p.paragraph_format
        pf.left_indent = Cm(0.5)
        pf.first_line_indent = Cm(-0.5)
        run = p.add_run(ref)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10)

    output_path = HERE / 'jurnal.docx'
    doc.save(str(output_path))
    print(f'✅ Jurnal berhasil digenerate: {output_path}')


if __name__ == '__main__':
    build_jurnal()

import sys, re, subprocess
from pathlib import Path
from build_skripsi import (
    SKRIPSI_DIR, TEMPLATE_PATH,
    parse_markdown, add_paragraph_with_italic, add_code_block,
    add_table_from_lines,
)
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ── Numbering ──

def set_num(para, num_id, ilvl):
    pass


def make_chapter_num(doc, abs_id, ch):
    return None


def find_abs_id(doc):
    for a in doc.part.numbering_part.element.findall(qn('w:abstractNum')):
        for l in a.findall(qn('w:lvl')):
            lt = l.find(qn('w:lvlText'))
            if lt is not None and lt.get(qn('w:val')) == '%1.%2.':
                return int(a.get(qn('w:abstractNumId')))
    return 0


# ── Page Setup ──

def copy_margins(t, s):
    for a in ['page_width','page_height','top_margin','bottom_margin','left_margin','right_margin']:
        setattr(t, a, getattr(s, a))


def add_body_sec(doc):
    sec = doc.add_section()
    copy_margins(sec, doc.sections[0])
    sp = sec._sectPr
    ex = sp.find(qn('w:pgNumType'))
    if ex is not None: sp.remove(ex)
    pn = OxmlElement('w:pgNumType'); pn.set(qn('w:fmt'), 'decimal'); sp.append(pn)
    sec.different_first_page_header_footer = True

    fp = sec.first_page_footer
    try: fp.is_linked_to_previous = False
    except: pass
    pp = fp.paragraphs[0] if fp.paragraphs else fp.add_paragraph()
    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER; pp.paragraph_format.first_line_indent = Cm(0)
    r0 = pp.add_run(); f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin'); r0._r.append(f1)
    r1 = pp.add_run(); r1.font.name='Times New Roman'; r1.font.size=Pt(12)
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text=' PAGE '; r1._r.append(it)
    r2 = pp.add_run(); r2.font.name='Times New Roman'; r2.font.size=Pt(12)
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end'); r2._r.append(f2)

    df = sec.footer
    try: df.is_linked_to_previous = False
    except: pass
    for p in df.paragraphs: p.clear()

    dh = sec.header
    try: dh.is_linked_to_previous = False
    except: pass
    hp = dh.paragraphs[0] if dh.paragraphs else dh.add_paragraph()
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT; hp.paragraph_format.first_line_indent = Cm(0)
    r0 = hp.add_run(); f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin'); r0._r.append(f1)
    r1 = hp.add_run(); r1.font.name='Times New Roman'; r1.font.size=Pt(10)
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text=' PAGE '; r1._r.append(it)
    r2 = hp.add_run(); r2.font.name='Times New Roman'; r2.font.size=Pt(10)
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end'); r2._r.append(f2)

    fh = sec.first_page_header
    try: fh.is_linked_to_previous = False
    except: pass
    for p in fh.paragraphs: p.clear()


# ── TOC ──

def inject_toc(doc, code):
    p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.left_indent = Cm(0)
    r0 = p.add_run(); f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin'); r0._r.append(f1)
    r1 = p.add_run(); it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = code; r1._r.append(it)
    r2 = p.add_run(); f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'separate'); r2._r.append(f2)
    r3 = p.add_run('(Update)'); r3.font.size = Pt(10); r3.font.italic = True
    r4 = p.add_run(); f3 = OxmlElement('w:fldChar'); f3.set(qn('w:fldCharType'), 'end'); r4._r.append(f3)


def update_lo(path):
    try:
        subprocess.run(['libreoffice','--headless','--convert-to','docx',str(path),
                        '--outdir',str(path.parent)], capture_output=True, timeout=30)
    except: pass


# ── Captions ──

def add_seq(para, seq_name, counter):
    r0 = para.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin'); r0._r.append(f1)
    r1 = para.add_run()
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = f' SEQ {seq_name} \\* ARABIC '; r1._r.append(it)
    r2 = para.add_run()
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'separate'); r2._r.append(f2)
    r3 = para.add_run(str(counter)); r3.font.size = Pt(10); r3.font.bold = False
    r4 = para.add_run()
    f3 = OxmlElement('w:fldChar'); f3.set(qn('w:fldCharType'), 'end'); r4._r.append(f3)


def set_caption_style(doc):
    cap = doc.styles['Caption']
    cap.font.name = 'Times New Roman'; cap.font.size = Pt(10); cap.font.italic = False; cap.font.bold = False
    cap.paragraph_format.line_spacing = 1.0; cap.paragraph_format.first_line_indent = Cm(0)
    cap.paragraph_format.space_before = Pt(0); cap.paragraph_format.space_after = Pt(0)


def add_img_cap(doc, chapter, counter, text):
    set_caption_style(doc)
    p = doc.add_paragraph(style='Caption')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0)
    p.clear()
    run = p.add_run(f'{chapter}.')
    run.font.name = 'Times New Roman'; run.font.size = Pt(10); run.font.italic = False; run.font.bold = False
    add_seq(p, 'Figure', counter)
    run2 = p.add_run(f' {text}')
    run2.font.name = 'Times New Roman'; run2.font.size = Pt(10); run2.font.italic = False; run2.font.bold = False


def add_tbl_cap(doc, chapter, counter, text):
    set_caption_style(doc)
    p = doc.add_paragraph(style='Caption')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0)
    p.clear()
    run = p.add_run(f'{chapter}.')
    run.font.name = 'Times New Roman'; run.font.size = Pt(10); run.font.italic = False; run.font.bold = False
    add_seq(p, 'Table', counter)
    run2 = p.add_run(f' {text}')
    run2.font.name = 'Times New Roman'; run2.font.size = Pt(10); run2.font.italic = False; run2.font.bold = False


# ── Main ──

tp = SKRIPSI_DIR / 'template-skripsi.docx'
mp = SKRIPSI_DIR / 'test-formatting.md'
op = SKRIPSI_DIR / 'test-output.docx'

if not tp.exists():
    print('Template tidak ditemukan. Jalankan: python create_template.py')
    sys.exit(1)

doc = Document(str(tp))
for p in doc.paragraphs: p.clear()

abs_id = find_abs_id(doc)
blocks = parse_markdown(mp)
print(f'Memproses ({len(blocks)} blok)')

# ── Preliminary ──
doc.add_heading('DAFTAR ISI', level=1)
inject_toc(doc, 'TOC \\o "1-3" \\h \\z \\u')
doc.add_page_break()
doc.add_heading('DAFTAR GAMBAR', level=1)
inject_toc(doc, 'TOC \\c "Figure" \\h \\z \\u')
doc.add_page_break()
doc.add_heading('DAFTAR TABEL', level=1)
inject_toc(doc, 'TOC \\c "Table" \\h \\z \\u')
doc.add_page_break()

# ── Body ──
ch = 0
cid = None
fc = 0
tc = 0

for block in blocks:
    t = block['type']
    if t == 'chapter':
        ch += 1; fc = 0; tc = 0
        add_body_sec(doc)
        cid = make_chapter_num(doc, abs_id, ch)
        roman = {1:'I',2:'II',3:'III',4:'IV',5:'V',6:'VI',7:'VII',8:'VIII',9:'IX',10:'X'}.get(ch, str(ch))
        h1 = doc.add_paragraph(f'BAB {roman}', style='Heading 1'); h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ct = doc.add_paragraph(block['title'], style='ChapterTitle'); ct.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ep = doc.add_paragraph(''); ep.paragraph_format.first_line_indent = Cm(0); ep.paragraph_format.line_spacing = 1.0
        print(f'  BAB {roman}: {block["title"]}')

    elif t == 'heading':
        lv = block['level']; tx = block['text']
        clean = re.sub(r'^[\d\.]+\s+', '', tx)
        if lv == 1:
            p = doc.add_paragraph(tx, style='Heading 1'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif lv == 2:
            doc.add_paragraph('', style='Normal').paragraph_format.first_line_indent = Cm(0)
            sp = tx.split(' ', 1)
            text = sp[0] + '	' + sp[1] if len(sp) > 1 else tx
            p = doc.add_paragraph(text, style='Heading 2')
        elif lv == 3:
            doc.add_paragraph('', style='Normal').paragraph_format.first_line_indent = Cm(0)
            sp = tx.split(' ', 1)
            text = sp[0] + '	' + sp[1] if len(sp) > 1 else tx
            p = doc.add_paragraph(text, style='Heading 3')

    elif t == 'text':
        tx = block['content'].strip()
        if tx:
            m = re.match(r'^!\[(.*)\]\((.+)\)$', tx)
            if m:
                fc += 1
                add_img_cap(doc, ch, fc, m.group(1))
                print(f'  Figure: {ch}.{fc}')
            else:
                add_paragraph_with_italic(doc, tx)

    elif t == 'code':
        if block['content'].strip(): add_code_block(doc, block['content'])

    elif t == 'table':
        tc += 1
        add_tbl_cap(doc, ch, tc, f'Data perbandingan')
        add_table_from_lines(doc, block['content'])
        print(f'  Table: {ch}.{tc}')

doc.save(str(op))
print(f'\nSaved: {op}')
print('Buka file di Word → Ctrl+A → F9 untuk update Daftar Isi.')
